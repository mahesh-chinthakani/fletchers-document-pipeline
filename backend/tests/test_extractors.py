from io import BytesIO

import pytest
from docx import Document

from backend.app.extractors.docx import DocxExtractor
from backend.app.extractors.pdf import PdfExtractor
from backend.app.extractors.registry import ExtractorRegistry
from backend.app.services.ocr import MockOCRService


def test_registry_returns_pdf_extractor():
    registry = ExtractorRegistry()

    extractor = registry.get_extractor("example.pdf")

    assert isinstance(extractor, PdfExtractor)


def test_registry_returns_docx_extractor():
    registry = ExtractorRegistry()

    extractor = registry.get_extractor("example.docx")

    assert isinstance(extractor, DocxExtractor)


def test_registry_rejects_unsupported_file_type():
    registry = ExtractorRegistry()

    with pytest.raises(ValueError, match="Unsupported document type"):
        registry.get_extractor("example.txt")


def test_docx_extractor_returns_non_empty_paragraphs():
    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph")

    buffer = BytesIO()
    document.save(buffer)

    extractor = DocxExtractor()
    result = extractor.extract(buffer.getvalue())

    assert result.extraction_method == "docx"
    assert len(result.chunks) == 2
    assert result.chunks[0].content == "First paragraph"
    assert result.chunks[1].content == "Second paragraph"


def test_pdf_with_text_uses_digital_extraction(monkeypatch):
    class FakePage:
        def extract_text(self):
            return "This is enough digitally extractable text for the test."

    class FakeReader:
        def __init__(self, _):
            self.pages = [FakePage()]

    monkeypatch.setattr(
        "backend.app.extractors.pdf.PdfReader",
        FakeReader,
    )

    extractor = PdfExtractor(MockOCRService())
    result = extractor.extract(b"fake-pdf-bytes")

    assert result.extraction_method == "digital_pdf"
    assert len(result.chunks) == 1
    assert result.chunks[0].page_number == 1


def test_pdf_without_text_routes_to_ocr(monkeypatch):
    class FakePage:
        def extract_text(self):
            return ""

    class FakeReader:
        def __init__(self, _):
            self.pages = [FakePage()]

    monkeypatch.setattr(
        "backend.app.extractors.pdf.PdfReader",
        FakeReader,
    )

    extractor = PdfExtractor(MockOCRService())
    result = extractor.extract(b"fake-pdf-bytes")

    assert result.extraction_method == "mock_ocr"
    assert len(result.chunks) == 1
    assert "OCR extraction was required" in result.chunks[0].content